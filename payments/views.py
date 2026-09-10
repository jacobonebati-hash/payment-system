import os

from decimal import Decimal

from django.contrib import messages

from django.http import (
    FileResponse,
    HttpResponse,
)

from django.shortcuts import (
    render,
    redirect,
    get_object_or_404,
)

from openpyxl import Workbook

from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
)

from .models import Member

from .forms import (
    MemberForm,
    PaymentForm,
    ExcelImportForm,
)

from .excel import generate_member_excel

from .excel_importer import (
    import_excel_file,
)


def dashboard(request):

    members = Member.objects.all()

    total_required = Decimal("0.00")

    total_paid = Decimal("0.00")

    total_debt = Decimal("0.00")

    completed_members = 0

    debt_members = 0

    for member in members:

        required = (
            member.kiasi_anachotakiwa
        )

        paid = (
            member.jumla_ya_malipo()
        )

        debt = (
            member.deni()
        )

        total_required += required

        total_paid += paid

        total_debt += debt

        if debt <= 0:

            completed_members += 1

        else:

            debt_members += 1

    context = {

        "members": members,

        "total_required": total_required,

        "total_paid": total_paid,

        "total_debt": total_debt,

        "total_members": members.count(),

        "completed_members": completed_members,

        "debt_members": debt_members,
    }

    return render(
        request,
        "payments/dashboard.html",
        context
    )


def member_list(request):

    members = Member.objects.all().order_by(
        "jina"
    )

    return render(
        request,
        "payments/member_list.html",
        {
            "members": members
        }
    )


def member_create(request):

    if request.method == "POST":

        form = MemberForm(
            request.POST
        )

        if form.is_valid():

            member = form.save()

            generate_member_excel(
                member
            )

            messages.success(
                request,
                "Mwanachama ameongezwa."
            )

            return redirect(
                "member_list"
            )

    else:

        form = MemberForm()

    return render(
        request,
        "payments/member_form.html",
        {
            "form": form
        }
    )


def member_detail(
    request,
    pk
):

    member = get_object_or_404(
        Member,
        pk=pk
    )

    payments = member.payments.all().order_by(
        "-tarehe"
    )

    return render(
        request,
        "payments/member_detail.html",
        {
            "member": member,
            "payments": payments,
        }
    )


def payment_create(request):

    if request.method == "POST":

        form = PaymentForm(
            request.POST
        )

        if form.is_valid():

            payment = form.save()

            messages.success(
                request,
                "Malipo yamehifadhiwa."
            )

            return redirect(
                "member_detail",
                pk=payment.member.id
            )

    else:

        form = PaymentForm()

    return render(
        request,
        "payments/payment_form.html",
        {
            "form": form
        }
    )


def import_excel(request):

    if request.method == "POST":

        form = ExcelImportForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            excel_file = form.cleaned_data[
                "excel_file"
            ]

            required_amount = form.cleaned_data[
                "kiasi_anachotakiwa"
            ]

            try:

                result = import_excel_file(

                    excel_file,

                    required_amount
                )

                if result["errors"]:

                    messages.warning(

                        request,

                        (
                            "Excel imeingizwa lakini "
                            f"kuna errors "
                            f"{len(result['errors'])}."
                        )
                    )

                else:

                    messages.success(

                        request,

                        (
                            "Excel imefanikiwa "
                            "kuingizwa. "
                            f"Walioongezwa: "
                            f"{result['members']}. "
                            f"Malipo: "
                            f"{result['payments']}. "
                            f"Rows zilizorukwa: "
                            f"{result['skipped']}."
                        )
                    )

                return redirect(
                    "member_list"
                )

            except Exception as error:

                messages.error(

                    request,

                    f"Kuna tatizo: {error}"
                )

    else:

        form = ExcelImportForm()

    return render(

        request,

        "payments/excel_import.html",

        {
            "form": form
        }
    )


def download_member_excel(
    request,
    pk
):

    member = get_object_or_404(
        Member,
        pk=pk
    )

    file_path = generate_member_excel(
        member
    )

    return FileResponse(

        open(
            file_path,
            "rb"
        ),

        as_attachment=True,

        filename=os.path.basename(
            file_path
        )
    )


def export_excel(request):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Wanachama"

    headers = [

        "Na.",

        "Jina",

        "Simu",

        "Eneo",

        "Anachotakiwa",

        "Jumla Alicholipa",

        "Kilichobaki",

        "Status",
    ]

    for col, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=col
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    members = Member.objects.all().order_by(
        "jina"
    )

    for index, member in enumerate(
        members,
        start=1
    ):

        row = index + 1

        sheet.cell(
            row=row,
            column=1
        ).value = index

        sheet.cell(
            row=row,
            column=2
        ).value = member.jina

        sheet.cell(
            row=row,
            column=3
        ).value = member.simu

        sheet.cell(
            row=row,
            column=4
        ).value = member.eneo

        sheet.cell(
            row=row,
            column=5
        ).value = float(
            member.kiasi_anachotakiwa
        )

        sheet.cell(
            row=row,
            column=6
        ).value = float(
            member.jumla_ya_malipo()
        )

        sheet.cell(
            row=row,
            column=7
        ).value = float(
            member.deni()
        )

        sheet.cell(
            row=row,
            column=8
        ).value = member.status()

    widths = {

        "A": 8,

        "B": 30,

        "C": 20,

        "D": 25,

        "E": 20,

        "F": 20,

        "G": 20,

        "H": 18,
    }

    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width

    sheet.freeze_panes = "A2"

    response = HttpResponse(

        content_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="wanachama_malipo.xlsx"'
    )

    workbook.save(response)

    return response