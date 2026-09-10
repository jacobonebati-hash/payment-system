from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from django.db import transaction

from .models import Member, Payment


# =========================================================
# 1. KUSAFISHA NAMBA YA SIMU
# =========================================================

def clean_phone(value):
    if value is None:
        return ""

    phone = str(value).strip()

    # Ondoa apostrophe, space, dash na mabano
    phone = (
        phone
        .replace("'", "")
        .replace(" ", "")
        .replace("-", "")
        .replace("(", "")
        .replace(")", "")
    )

    # +255XXXXXXXXX -> 255XXXXXXXXX
    if phone.startswith("+255"):
        phone = phone[1:]

    # 07XXXXXXXX / 06XXXXXXXX -> 2557XXXXXXXX / 2556XXXXXXXX
    elif phone.startswith("0") and len(phone) == 10:
        phone = "255" + phone[1:]

    return phone


def is_phone(value):
    """
    Tanzania phone number:
    2556XXXXXXXX
    2557XXXXXXXX
    """

    phone = clean_phone(value)

    if not phone.isdigit():
        return False

    if len(phone) != 12:
        return False

    if not phone.startswith("255"):
        return False

    # Tanzania mobile numbers huanza 6 au 7
    return phone[3] in ["6", "7"]


# =========================================================
# 2. KUSOMA KIASI CHA FEDHA
# =========================================================

def convert_amount(value):
    if value is None:
        return Decimal("0.00")

    if isinstance(value, bool):
        return Decimal("0.00")

    if isinstance(value, (int, float)):
        try:
            return Decimal(str(value))
        except Exception:
            return Decimal("0.00")

    text = str(value).strip()

    if not text:
        return Decimal("0.00")

    # Ondoa currency notation
    text = (
        text
        .replace(",", "")
        .replace("TSh", "")
        .replace("TSH", "")
        .replace("TZS", "")
        .replace("tsh", "")
        .strip()
    )

    try:
        amount = Decimal(text)

        if amount < 0:
            return Decimal("0.00")

        return amount

    except (InvalidOperation, ValueError):
        return Decimal("0.00")


# =========================================================
# 3. KUTAMBUA HEADING
# =========================================================

def is_heading(value):
    if not isinstance(value, str):
        return False

    text = value.upper().strip()

    headings = [
        "JINA",
        "NAME",
        "TAARIFA",
        "AMBAO WAMEKIDHI",
        "WAMEKIDHI VIGEZO",
        "VYA UANACHAMA",
        "UANACHAMA",
        "WALIOLIPA",
        "JUMLA",
        "TOTAL",
        "SIMU",
        "PHONE",
        "NO.",
        "NA.",
    ]

    return any(word in text for word in headings)


# =========================================================
# 4. KUTAFUTA JINA
# =========================================================

def find_name(row):

    for value in row:

        if value is None:
            continue

        text = str(value).strip()

        if len(text) < 3:
            continue

        if is_heading(text):
            continue

        # Usichukue namba tupu kama jina
        if text.replace(" ", "").isdigit():
            continue

        # Usichukue simu kama jina
        if is_phone(text):
            continue

        return text

    return ""


# =========================================================
# 5. KUTAFUTA SIMU
# =========================================================

def find_phone(row):

    for value in row:

        if is_phone(value):
            return clean_phone(value)

    return ""


# =========================================================
# 6. KUTAMBUA TOTAL / JUMLA
# =========================================================

def looks_like_total(value):

    if value is None:
        return False

    if isinstance(value, str):

        text = value.upper().strip()

        total_words = [
            "TOTAL",
            "JUMLA",
            "JUMla",
            "JML",
        ]

        return any(word in text for word in total_words)

    return False


# =========================================================
# 7. KUSOMA MALIPO
# =========================================================

def find_payments(row):

    amounts = []

    for value in row:

        # Usichukue simu kama malipo
        if is_phone(value):
            continue

        # Text inayojulikana kuwa heading/total
        if looks_like_total(value):
            continue

        amount = convert_amount(value)

        # Malipo yetu yanaanzia 100
        if amount < Decimal("100"):
            continue

        # Zuia namba kubwa sana
        # Hii inazuia simu zilizoharibika kuingia
        if amount >= Decimal("1000000000"):
            continue

        amounts.append(amount)

    return amounts


# =========================================================
# 8. KUONDOA TOTAL IKIWA IMEWEKWA MWISHO
# =========================================================

def remove_duplicate_total(amounts):

    if len(amounts) < 2:
        return amounts

    total = amounts[-1]

    payments = amounts[:-1]

    calculated_total = sum(
        payments,
        Decimal("0.00")
    )

    # Mfano:
    # 5000 + 5000 + 10000
    #
    # Excel:
    # 5000 | 5000 | 10000
    #
    # 10000 ni total, hivyo iondolewe.

    if total == calculated_total:
        return payments

    return amounts


# =========================================================
# 9. IMPORT EXCEL
# =========================================================

@transaction.atomic
def import_excel_file(excel_file, default_required_amount):

    workbook = load_workbook(
        excel_file,
        data_only=True
    )

    sheet = workbook.active

    imported_members = 0
    imported_payments = 0
    skipped_rows = 0
    duplicate_payments = 0
    errors = []

    for row_number, row in enumerate(
        sheet.iter_rows(values_only=True),
        start=1
    ):

        try:

            # -----------------------------------------
            # Skip empty rows
            # -----------------------------------------

            if not any(
                value not in (None, "")
                for value in row
            ):
                continue

            # -----------------------------------------
            # Tafuta jina
            # -----------------------------------------

            name = find_name(row)

            if not name:
                skipped_rows += 1
                continue

            # -----------------------------------------
            # Tafuta simu
            # -----------------------------------------

            phone = find_phone(row)

            # -----------------------------------------
            # Tafuta mwanachama
            # -----------------------------------------

            member = None

            if phone:

                member = Member.objects.filter(
                    simu=phone
                ).first()

            if member is None:

                member = Member.objects.filter(
                    jina__iexact=name
                ).first()

            # -----------------------------------------
            # Create member
            # -----------------------------------------

            if member is None:

                member = Member.objects.create(
                    jina=name,
                    simu=phone,
                    kiasi_anachotakiwa=default_required_amount,
                )

                imported_members += 1

            else:

                changed = False

                if phone and not member.simu:

                    member.simu = phone
                    changed = True

                if changed:

                    member.save(
                        update_fields=["simu"]
                    )

            # -----------------------------------------
            # Soma payments
            # -----------------------------------------

            amounts = find_payments(row)

            # Ondoa total kama ipo
            amounts = remove_duplicate_total(amounts)

            # -----------------------------------------
            # Save payments
            # -----------------------------------------

            for amount in amounts:

                # =====================================
                # SAFETY CHECK
                # =====================================

                if amount >= Decimal("1000000000"):
                    continue

                if amount <= Decimal("0"):
                    continue

                # =====================================
                # DUPLICATE PROTECTION
                # =====================================

                possible_duplicate = Payment.objects.filter(
                    member=member,
                    kiasi=amount,
                    maelezo=f"Imported kutoka Excel - row {row_number}",
                ).exists()

                if possible_duplicate:

                    duplicate_payments += 1
                    continue

                Payment.objects.create(

                    member=member,

                    kiasi=amount,

                    payment_method="Nyingine",

                    maelezo=(
                        f"Imported kutoka Excel "
                        f"- row {row_number}"
                    ),
                )

                imported_payments += 1

        except Exception as error:

            errors.append(
                f"Row {row_number}: {error}"
            )

    return {

        "members": imported_members,

        "payments": imported_payments,

        "duplicates": duplicate_payments,

        "skipped": skipped_rows,

        "errors": errors,
    }