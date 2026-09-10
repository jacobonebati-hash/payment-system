import os

from django.conf import settings

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .models import Member


def generate_member_excel(member):

    folder = os.path.join(
        settings.BASE_DIR,
        "excel_reports"
    )

    os.makedirs(
        folder,
        exist_ok=True
    )

    safe_name = "".join(
        c
        for c in member.jina
        if c.isalnum() or c in (" ", "_", "-")
    ).strip()

    safe_name = safe_name.replace(
        " ",
        "_"
    )

    file_path = os.path.join(
        folder,
        f"{safe_name}_{member.id}.xlsx"
    )

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Taarifa za Malipo"

    # =========================
    # TITLE
    # =========================

    sheet.merge_cells("A1:H1")

    sheet["A1"] = "TAARIFA ZA MALIPO"

    sheet["A1"].font = Font(
        bold=True,
        size=18
    )

    sheet["A1"].alignment = Alignment(
        horizontal="center"
    )

    # =========================
    # MEMBER INFORMATION
    # =========================

    taarifa = [
        ("Jina", member.jina),
        ("Simu", member.simu),
        ("Eneo", member.eneo),
        (
            "Anachotakiwa",
            float(member.kiasi_anachotakiwa)
        ),
        (
            "Jumla Alicholipa",
            float(member.jumla_ya_malipo())
        ),
        (
            "Kilichobaki",
            float(member.deni())
        ),
        (
            "Status",
            member.status()
        ),
    ]

    row = 3

    for label, value in taarifa:

        sheet[f"A{row}"] = label

        sheet[f"B{row}"] = value

        sheet[f"A{row}"].font = Font(
            bold=True
        )

        row += 1

    # =========================
    # PAYMENT TABLE
    # =========================

    start_row = 12

    headers = [
        "Na.",
        "Tarehe",
        "Kiasi",
        "Njia ya Malipo",
        "Transaction / Risiti",
        "Maelezo",
        "Jumla Iliyolipwa",
        "Kilichobaki",
    ]

    for col, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=start_row,
            column=col
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # =========================
    # PAYMENTS
    # =========================

    payments = (
        member.payments
        .all()
        .order_by("tarehe")
    )

    running_total = 0

    row = start_row + 1

    for index, payment in enumerate(
        payments,
        start=1
    ):

        running_total += float(
            payment.kiasi
        )

        remaining = (
            float(member.kiasi_anachotakiwa)
            - running_total
        )

        if remaining < 0:
            remaining = 0

        sheet.cell(
            row=row,
            column=1
        ).value = index

        sheet.cell(
            row=row,
            column=2
        ).value = payment.tarehe.strftime(
            "%d/%m/%Y %H:%M"
        )

        sheet.cell(
            row=row,
            column=3
        ).value = float(
            payment.kiasi
        )

        sheet.cell(
            row=row,
            column=4
        ).value = (
            payment
            .get_payment_method_display()
        )

        sheet.cell(
            row=row,
            column=5
        ).value = payment.transaction_number

        sheet.cell(
            row=row,
            column=6
        ).value = payment.maelezo

        sheet.cell(
            row=row,
            column=7
        ).value = running_total

        sheet.cell(
            row=row,
            column=8
        ).value = remaining

        row += 1

    # =========================
    # TOTAL
    # =========================

    total_row = row + 1

    sheet.cell(
        row=total_row,
        column=1
    ).value = "JUMLA"

    sheet.cell(
        row=total_row,
        column=1
    ).font = Font(
        bold=True
    )

    sheet.cell(
        row=total_row,
        column=3
    ).value = float(
        member.jumla_ya_malipo()
    )

    sheet.cell(
        row=total_row,
        column=3
    ).font = Font(
        bold=True
    )

    sheet.cell(
        row=total_row,
        column=8
    ).value = float(
        member.deni()
    )

    sheet.cell(
        row=total_row,
        column=8
    ).font = Font(
        bold=True
    )

    # =========================
    # BORDERS
    # =========================

    thin = Side(
        style="thin",
        color="B7B7B7"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    for row_cells in sheet.iter_rows(
        min_row=start_row,
        max_row=total_row,
        min_col=1,
        max_col=8
    ):

        for cell in row_cells:

            cell.border = border

    # =========================
    # NUMBER FORMAT
    # =========================

    for r in range(
        start_row + 1,
        total_row + 1
    ):

        for c in [3, 7, 8]:

            sheet.cell(
                row=r,
                column=c
            ).number_format = "#,##0.00"

    # =========================
    # COLUMN WIDTHS
    # =========================

    widths = {
        "A": 8,
        "B": 20,
        "C": 18,
        "D": 20,
        "E": 25,
        "F": 30,
        "G": 20,
        "H": 20,
    }

    for column, width in widths.items():

        sheet.column_dimensions[
            column
        ].width = width

    sheet.freeze_panes = "A13"

    workbook.save(
        file_path
    )

    return file_path